from io import BytesIO

import pymupdf
from docx import Document


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    lowered_filename = filename.lower()

    if lowered_filename.endswith(".pdf"):
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document)
    elif lowered_filename.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraphs.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        text = "\n".join(paragraphs)
    else:
        raise ValueError("Only PDF and DOCX resumes are supported.")

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the uploaded resume.")

    return text
